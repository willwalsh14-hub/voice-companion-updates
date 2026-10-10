"""One release outline generates all accessible release-note formats.

Reference structure: title (H1), New Apps and Features / Fixes and
Enhancements (H2), numbered topics, lettered details, Roman subdetails.
Never copy product-specific HIMS content into Voice Companion documentation.
"""
import html
import json
from pathlib import Path
import re
import uuid
from xml.etree import ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED
from braille_documents import build_brf

class Verbatim(str):
    def title(self):
        return str(self)

STEM = 'Voice Companion Release Notes'
FORMATS = ('.txt', '.html', '.docx', '.epub', ' - DAISY 3.zip', '.brf')


def load_notes(root):
    root = Path(root)
    notes = json.loads((root / 'release-notes.json').read_text(encoding='utf-8'))
    version = re.search(r"APP_VERSION = '([^']+)'", (root / 'companion.py').read_text())[1]
    if notes['version'] != version:
        raise ValueError('Write release notes for this version before building or publishing.')
    if notes['status'] not in ('Testing build', 'Official release'):
        raise ValueError('Release status must be Testing build or Official release.')
    if ('test' in version.lower()) != (notes['status'] == 'Testing build'):
        raise ValueError('Release notes must distinguish testing builds from official releases.')
    if [s['title'] for s in notes['sections']] != ['New Apps and Features', 'Fixes and Enhancements']:
        raise ValueError('Release sections must follow the reference heading structure.')
    for section in notes['sections']:
        if not section['topics']:
            raise ValueError('Both release sections require reviewed content.')
        for topic in section['topics']:
            if not topic['title'] or not topic['details'] or len(topic['details']) > 26:
                raise ValueError('Each numbered topic needs a title and one to 26 lettered details.')
    return notes


def blocks(notes):
    output = [(None, [f"Version: {notes['version']}. {notes['status']}. Released {notes['date']}.",
                       f"Changes since {notes['baseline']}. {notes['scope']}"])]
    for section in notes['sections']:
        paragraphs = []
        for i, topic in enumerate(section['topics'], 1):
            paragraphs.append(f"{i}. {topic['title']}")
            for j, detail in enumerate(topic['details']):
                paragraphs.append(f"{chr(97+j)}. {detail}")
        output.append((Verbatim(section['title']), paragraphs))
    return output


def markup(items):
    result = []
    for n, (heading, paragraphs) in enumerate(items):
        if not heading:
            result.extend('<p>'+html.escape(p)+'</p>' for p in paragraphs)
            continue
        result.append(f'<section aria-labelledby="section-{n}"><h2 id="section-{n}">{html.escape(heading)}</h2><ol>')
        opened = False
        for p in paragraphs:
            label, text = p.split('. ', 1)
            if label.isdigit():
                if opened: result.append('</ol></li>')
                result.append('<li>'+html.escape(text)+'<ol type="a">')
                opened = True
            else:
                result.append('<li>'+html.escape(text)+'</li>')
        if opened: result.append('</ol></li>')
        result.append('</ol></section>')
    return '\n'.join(result)


def word(root, title, notes):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(.75)
    for name, size in [('Normal',11),('Heading 1',20),('Heading 2',15)]:
        style = doc.styles[name]
        style.font.name = 'Aptos'
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0,0,0)
    doc.styles['Normal'].paragraph_format.space_after = Pt(5)
    doc.add_heading(title, 1)
    for p in blocks(notes)[0][1]: doc.add_paragraph(p)
    numbering = doc.part.numbering_part.element
    abstract_id = max(int(e.get(qn('w:abstractNumId'))) for e in numbering.findall(qn('w:abstractNum'))) + 1
    abstract = OxmlElement('w:abstractNum');abstract.set(qn('w:abstractNumId'),str(abstract_id))
    for level, fmt in enumerate(('decimal','lowerLetter','lowerRoman')):
        node=OxmlElement('w:lvl');node.set(qn('w:ilvl'),str(level))
        for tag,val in [('start','1'),('numFmt',fmt),('lvlText',f'%{level+1}.'),('lvlJc','left')]:
            child=OxmlElement('w:'+tag);child.set(qn('w:val'),val);node.append(child)
        props=OxmlElement('w:pPr');indent=OxmlElement('w:ind')
        indent.set(qn('w:left'),str(360*(level+1)));indent.set(qn('w:hanging'),'360');props.append(indent);node.append(props)
        abstract.append(node)
    numbering.append(abstract)
    for s in notes['sections']:
        doc.add_heading(s['title'],2)
        num=numbering.add_num(abstract_id)
        num.add_lvlOverride(0).add_startOverride(1)
        numid=num.get(qn('w:numId'))
        for topic in s['topics']:
            for level,text in [(0,topic['title'])]+[(1,d) for d in topic['details']]:
                p=doc.add_paragraph(text)
                props=p._p.get_or_add_pPr();np=OxmlElement('w:numPr')
                for tag,val in [('ilvl',str(level)),('numId',numid)]:
                    child=OxmlElement('w:'+tag);child.set(qn('w:val'),val);np.append(child)
                props.append(np)
                if level==0:p.paragraph_format.keep_with_next=True
    doc.core_properties.title=title;doc.core_properties.language='en-US'
    doc.save(root/(STEM+'.docx'))


def daisy(root, title, notes, uid):
    esc=html.escape
    nodes=[('title',title)]
    body=[f'<level1 id="release"><h1 id="title" smilref="guide.smil#text-title">{esc(title)}</h1>']
    for i,p in enumerate(blocks(notes)[0][1]):
        key=f'intro-{i}';nodes.append((key,p));body.append(f'<p id="{key}" smilref="guide.smil#text-{key}">{esc(p)}</p>')
    nav=[]
    for n,s in enumerate(notes['sections']):
        key=f'section-{n}';nodes.append((key,s['title']))
        body.append(f'<level2><h2 id="{key}" smilref="guide.smil#text-{key}">{esc(s["title"])}</h2><list type="ol" enum="1">')
        for i,t in enumerate(s['topics']):
            key=f'topic-{n}-{i}';nodes.append((key,t['title']))
            body.append(f'<li><p id="{key}" smilref="guide.smil#text-{key}">{esc(t["title"])}</p><list type="ol" enum="a">')
            for j,d in enumerate(t['details']):
                key=f'detail-{n}-{i}-{j}';nodes.append((key,d))
                body.append(f'<li><p id="{key}" smilref="guide.smil#text-{key}">{esc(d)}</p></li>')
            body.append('</list></li>')
        body.append('</list></level2>')
        nav.append(f'<navPoint id="nav-{n}" playOrder="{n+2}"><navLabel><text>{esc(s["title"])}</text></navLabel><content src="guide.smil#par-section-{n}"/></navPoint>')
    body.append('</level1>')
    dt=f'''<?xml version="1.0" encoding="utf-8"?><dtbook xmlns="http://www.daisy.org/z3986/2005/dtbook/" version="2005-3" xml:lang="en"><head><meta name="dc:Identifier" content="{uid}"/><meta name="dc:Title" content="{esc(title)}"/></head><book><frontmatter><doctitle>{esc(title)}</doctitle></frontmatter><bodymatter>{''.join(body)}</bodymatter></book></dtbook>'''
    pars=''.join(f'<par id="par-{k}"><text id="text-{k}" src="guide.xml#{k}"/></par>' for k,_ in nodes)
    smil=f'''<?xml version="1.0" encoding="utf-8"?><smil xmlns="http://www.w3.org/2001/SMIL20/"><head><meta name="dtb:uid" content="{uid}"/></head><body><seq id="all">{pars}</seq></body></smil>'''
    ncx=f'''<?xml version="1.0" encoding="utf-8"?><ncx xmlns="http://www.daisy.org/z3986/2005/ncx/" version="2005-1" xml:lang="en"><head><meta name="dtb:uid" content="{uid}"/><meta name="dtb:depth" content="2"/><meta name="dtb:totalPageCount" content="0"/><meta name="dtb:maxPageNumber" content="0"/></head><docTitle><text>{esc(title)}</text></docTitle><navMap><navPoint id="release" playOrder="1"><navLabel><text>{esc(title)}</text></navLabel><content src="guide.smil#par-title"/>{''.join(nav)}</navPoint></navMap></ncx>'''
    opf=f'''<?xml version="1.0" encoding="utf-8"?><package xmlns="http://openebook.org/namespaces/oeb-package/1.0/" unique-identifier="uid"><metadata><dc-metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:Title>{esc(title)}</dc:Title><dc:Identifier id="uid">{uid}</dc:Identifier><dc:Language>en</dc:Language></dc-metadata><x-metadata><meta name="dtb:multimediaType" content="textNCX"/><meta name="dtb:multimediaContent" content="text"/></x-metadata></metadata><manifest><item id="ncx" href="guide.ncx" media-type="application/x-dtbncx+xml"/><item id="dtbook" href="guide.xml" media-type="application/x-dtbook+xml"/><item id="smil" href="guide.smil" media-type="application/smil"/></manifest><spine toc="ncx"><itemref idref="smil"/></spine></package>'''
    with ZipFile(root/(STEM+' - DAISY 3.zip'),'w',ZIP_DEFLATED) as z:
        for name,text in [('guide.xml',dt),('guide.smil',smil),('guide.ncx',ncx),('guide.opf',opf),('READ ME.txt','Text-only DAISY 3. Extract all files together and open guide.opf in a DAISY 3 player supporting text-to-speech. No recorded narration is included.')]:z.writestr(name,text.encode('utf-8'))


def build(root):
    import build_user_guides as guides
    root=Path(root);notes=load_notes(root)
    title=Verbatim('Release Notes for Voice Companion '+notes['version'])
    items=blocks(notes)
    text=title+'\n\n'+'\n\n'.join(('\n\n'.join(([heading] if heading else [])+paragraphs)) for heading,paragraphs in items)+'\n'
    (root/(STEM+'.txt')).write_text(text,encoding='utf-8')
    saved=(guides.ROOT,guides.STEM,guides.UID,guides.content_markup)
    uid='urn:uuid:'+str(uuid.uuid5(uuid.NAMESPACE_URL,'voice-companion-release-'+notes['version']))
    try:
        guides.ROOT=root;guides.STEM=STEM;guides.UID=uid;guides.content_markup=markup
        # The existing builders title-case titles, which would corrupt case in commands.
        # All supplied release headings already use the required capitalization.
        guides.build_html(title,items);guides.build_epub(title,items)
    finally:
        guides.ROOT,guides.STEM,guides.UID,guides.content_markup=saved
    word(root,title,notes);daisy(root,title,notes,uid)
    build_brf(root/(STEM+'.txt'),root/(STEM+'.brf'))
    validate(root,notes)
    from guide_files import published_name
    import shutil
    for suffix in FORMATS:shutil.copyfile(root/(STEM+suffix),root/published_name(STEM+suffix,notes['version']))


def validate(root,notes=None):
    root=Path(root);notes=notes or load_notes(root)
    expected=[s['title'] for s in notes['sections']]
    with ZipFile(root/(STEM+'.docx')) as z:
        document=ET.fromstring(z.read('word/document.xml'))
    wn={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    paragraphs=document.findall('.//w:body/w:p',wn)
    texts=[''.join(e.text or '' for e in p.findall('.//w:t',wn)) for p in paragraphs]
    headings=[text for p,text in zip(paragraphs,texts)
              if p.find('w:pPr/w:pStyle',wn) is not None and
              p.find('w:pPr/w:pStyle',wn).get('{'+wn['w']+'}val')=='Heading2']
    if headings!=expected:raise ValueError('Word release outline differs.')
    page=(root/(STEM+'.html')).read_text(encoding='utf-8')
    if re.findall(r'<h2 id="section-\d+">([^<]+)</h2>',page)!=expected:
        raise ValueError('HTML release outline differs.')
    expected_text=[t['title'] for s in notes['sections'] for t in s['topics']]
    expected_details=[d for s in notes['sections'] for t in s['topics'] for d in t['details']]
    actual=texts
    if any(t not in actual for t in expected_text+expected_details):raise ValueError('Word release content is incomplete.')
    with ZipFile(root/(STEM+'.epub')) as z:
        page=ET.fromstring(z.read('OEBPS/content.xhtml'));ns={'x':'http://www.w3.org/1999/xhtml'}
        if [h.text for h in page.findall('.//x:h2',ns)]!=expected:raise ValueError('EPUB release outline differs.')
        details=[e.text for e in page.findall('.//x:ol/x:li/x:ol/x:li',ns)]
        if details!=expected_details:raise ValueError('EPUB lettered details differ.')
    with ZipFile(root/(STEM+' - DAISY 3.zip')) as z:
        dt=ET.fromstring(z.read('guide.xml'));ns={'d':'http://www.daisy.org/z3986/2005/dtbook/'}
        if [h.text for h in dt.findall('.//d:h2',ns)]!=expected:raise ValueError('DAISY release outline differs.')
        details=[e.text for e in dt.findall('.//d:list/d:li/d:list/d:li/d:p',ns)]
        if details!=expected_details:raise ValueError('DAISY lettered details differ.')
        ids={e.get('id') for e in dt.iter() if e.get('id')}
        smil=ET.fromstring(z.read('guide.smil'))
        if any(e.get('src').split('#')[1] not in ids for e in smil.iter() if e.tag.endswith('text')):raise ValueError('Broken DAISY navigation.')
    from braille_documents import validate_brf
    validate_brf((root/(STEM+'.brf')).read_bytes())
