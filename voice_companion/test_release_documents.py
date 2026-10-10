import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from docx import Document
from docx.oxml.ns import qn
import release_documents as release
from braille_documents import validate_brf, build_brf


def fixture(root):
    notes={'version':'1.2.3-test','status':'Testing build','date':'October 10, 2026','baseline':'1.2.2-test',
           'scope':'Fixture only.', 'sections':[
               {'title':'New Apps and Features','topics':[{'title':'First topic','details':['First detail','Second detail']},{'title':'Second topic','details':['Restart lettering']}]},
               {'title':'Fixes and Enhancements','topics':[{'title':'Restart numbering','details':['Final detail']}]}]}
    (root/'release-notes.json').write_text(json.dumps(notes))
    with patch.object(release,'build_brf',side_effect=lambda source,target:Path(target).write_bytes(b',FIXTURE\r\n')):
        release.build(root)
    return notes


class ReleaseDocumentTests(unittest.TestCase):
    def test_same_outline_content_and_numbering_in_all_formats(self):
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder);(root/'companion.py').write_text("APP_VERSION = '1.2.3-test'\n")
            notes=fixture(root);release.validate(root)
            text=(root/(release.STEM+'.txt')).read_text()
            self.assertIn('1. First topic\n\na. First detail\n\nb. Second detail',text)
            self.assertIn('2. Second topic\n\na. Restart lettering',text)
            self.assertIn('Fixes and Enhancements\n\n1. Restart numbering',text)
            doc=Document(root/(release.STEM+'.docx'))
            self.assertEqual([p.style.name for p in doc.paragraphs if p.style.name.startswith('Heading')],['Heading 1','Heading 2','Heading 2'])
            topics=[p for p in doc.paragraphs if p.text in ('First topic','Restart numbering')]
            ids=[p._p.find(qn('w:pPr')).find(qn('w:numPr')).find(qn('w:numId')).get(qn('w:val')) for p in topics]
            self.assertNotEqual(ids[0],ids[1])
            numbering=doc.part.numbering_part.element
            for numid in ids:
                number=next(n for n in numbering.findall(qn('w:num')) if n.get(qn('w:numId'))==numid)
                override=number.find(qn('w:lvlOverride'))
                self.assertEqual(override.get(qn('w:ilvl')),'0')
                self.assertEqual(override.find(qn('w:startOverride')).get(qn('w:val')),'1')
            with ZipFile(root/(release.STEM+'.epub')) as z:
                ns={'x':'http://www.w3.org/1999/xhtml'}
                tree=ET.fromstring(z.read('OEBPS/content.xhtml'))
                self.assertEqual(len(tree.findall('.//x:ol[@type="a"]',ns)),3)
                nav=ET.fromstring(z.read('OEBPS/nav.xhtml'))
                self.assertEqual([e.text for e in nav.findall('.//x:a',ns)],[s['title'] for s in notes['sections']])
                self.assertIn('1.2.3-test',tree.find('.//x:h1',ns).text)
            for suffix in release.FORMATS:self.assertTrue((root/(release.STEM+suffix)).is_file())

    def test_version_status_and_section_order_gate(self):
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder);(root/'companion.py').write_text("APP_VERSION = '1.2.3-test'\n")
            notes=fixture(root)
            for key,value in [('version','1.2.2-test'),('status','Official release')]:
                old=notes[key];notes[key]=value;(root/'release-notes.json').write_text(json.dumps(notes))
                with self.assertRaises(ValueError):release.load_notes(root)
                notes[key]=old
            notes['sections'].reverse();(root/'release-notes.json').write_text(json.dumps(notes))
            with self.assertRaises(ValueError):release.load_notes(root)

    def test_brf_rejects_print_unicode_and_page_overflow(self):
        for data in [b'',b'print lowercase',b'A'*41,b'\r\n'.join([b'A']*26),'\u2801'.encode('utf-8')]:
            with self.assertRaises(ValueError):validate_brf(data)
        validate_brf(b'A'*40+b'\r\n\f,A\r\n')

    def test_real_braille_is_contracted_and_back_translates(self):
        from braille_documents import translator
        import subprocess
        exe,env=translator()
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder);source=root/'input.txt';target=root/'output.brf'
            source.write_text('The quick brown fox\n')
            build_brf(source,target)
            self.assertEqual(target.read_bytes().strip(),b',! QK BR[N FOX')
            back=subprocess.run([exe,'-b','en-us-brf.dis,en-ueb-g2.ctb'],input=target.read_bytes().replace(b'\r\n',b'\n'),stdout=subprocess.PIPE,stderr=subprocess.PIPE,env=env,check=True)
            self.assertEqual(back.stdout.decode().strip(),'The quick brown fox')
