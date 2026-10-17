from pathlib import Path
import tempfile
import unittest
import zipfile
from document_reader import open_document

class DocumentReaderTests(unittest.TestCase):
    def test_text_unicode_html_and_word_are_read_without_changing_the_source(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder)
            text=root/'sample.txt';text.write_bytes('First line\nCafé'.encode('utf-16'))
            self.assertEqual(open_document(text).text,'First line\nCafé')
            html=root/'page.html';html.write_text('<h1>My heading</h1><p>Read this.</p><script>doNotRun()</script><style>hiddenCSS</style>')
            result=open_document(html);self.assertIn('My heading',result.text);self.assertNotIn('doNotRun',result.text)
            path=root/'sample.docx';doc=Document();doc.add_paragraph('Word paragraph');table=doc.add_table(1,2);table.cell(0,0).text='First cell';table.cell(0,1).text='Second cell';doc.save(path)
            before=path.read_bytes();result=open_document(path)
            self.assertIn('Word paragraph',result.text);self.assertIn('First cell',result.text);self.assertEqual(path.read_bytes(),before)
            result.reading.position=0;self.assertIn('Word',result.reading.move('word',0))
    def test_epub_spine_order_and_odt(self):
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder);epub=root/'book.epub'
            with zipfile.ZipFile(epub,'w') as z:
                z.writestr('META-INF/container.xml','<container><rootfiles><rootfile full-path="OPS/package.opf"/></rootfiles></container>')
                z.writestr('OPS/package.opf','<package xmlns="http://www.idpf.org/2007/opf"><manifest><item id="two" href="two.xhtml"/><item id="one" href="one.xhtml"/></manifest><spine><itemref idref="one"/><itemref idref="two"/></spine></package>')
                z.writestr('OPS/one.xhtml','<p>First chapter</p>');z.writestr('OPS/two.xhtml','<p>Second chapter</p>')
            text=open_document(epub).text;self.assertLess(text.index('First chapter'),text.index('Second chapter'))
            odt=root/'text.odt'
            with zipfile.ZipFile(odt,'w') as z:z.writestr('content.xml','<body xmlns:t="urn:test"><t:p>Open document text</t:p></body>')
            self.assertIn('Open document text',open_document(odt).text)
    def test_pdf_text_page_labels_scanned_and_password(self):
        from pypdf import PdfWriter
        from pypdf.generic import DictionaryObject,NameObject,DecodedStreamObject
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder);path=root/'text.pdf';writer=PdfWriter();page=writer.add_blank_page(300,300)
            font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
            page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):writer._add_object(font)})})
            stream=DecodedStreamObject();stream.set_data(b'BT /F1 12 Tf 20 250 Td (PDF readable text) Tj ET');page[NameObject('/Contents')]=writer._add_object(stream)
            with path.open('wb') as output:writer.write(output)
            text=open_document(path).text;self.assertIn('Page 1',text);self.assertIn('PDF readable text',text)
            blank=root/'scan.pdf';writer=PdfWriter();writer.add_blank_page(300,300)
            with blank.open('wb') as output:writer.write(output)
            self.assertIn('OCR is not included',open_document(blank).text)
            writer.encrypt('private');protected=root/'protected.pdf'
            with protected.open('wb') as output:writer.write(output)
            with self.assertRaisesRegex(ValueError,'password'):open_document(protected)
    def test_unsupported_and_oversized_archives_are_not_executed(self):
        from unittest.mock import patch
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'program.exe';path.write_text('Do not execute')
            with self.assertRaisesRegex(ValueError,'not supported'):open_document(path)
            path=Path(folder)/'bomb.docx'
            with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as archive:archive.writestr('word/document.xml','a'*2000)
            with patch('document_reader.MAX_BYTES',1000):
                with self.assertRaises(ValueError):open_document(path)
    def test_reader_caret_and_escape_dispatch(self):
        import companion as app
        from document_reader import DocumentReader
        from unittest.mock import patch,Mock
        reader=DocumentReader('sample.txt','First sentence.\nNext line.')
        with patch.object(app,'FILE_READER',reader),patch.object(app,'FILE_MANAGER',Mock(say_current=Mock(return_value='sample.txt, file.'))),patch.object(app,'APP_WINDOW',None),patch.object(app,'speak'),patch.object(app,'sync_app_context'):
            context=app.app_context('file_reader');self.assertTrue(context['readonly'])
            self.assertTrue(app.apply_keyboard_edit(dict(source=context['source'],after=reader.text,caret=5,seq=1,readonly=True),'file_reader'))
            self.assertEqual(reader.reading.position,5)
            self.assertEqual(app.handle('back to files','file_reader'),'files')

if __name__=='__main__':unittest.main()
