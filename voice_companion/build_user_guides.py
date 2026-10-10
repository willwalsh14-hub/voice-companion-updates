"""Generate accessible guide formats from START HERE - Veteran.txt.

Run after editing the TXT guide. The Windows installer builder runs this too.
"""
from __future__ import annotations

import html
import re
from datetime import datetime, timezone
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZIP_STORED, ZipFile

from docx import Document
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "START HERE - Veteran.txt"
STEM = "START HERE - Veteran"
UID = "urn:uuid:340725ad-11b0-4490-b546-9d4be0fa9e18"


def sections():
    lines = SOURCE.read_text(encoding="utf-8-sig").splitlines()
    title = lines[0].strip()
    assert title and title == title.upper()
    blocks = []
    current = None
    for line in lines[1:]:
        s = line.strip()
        if not s:
            continue
        if s == s.upper() and re.fullmatch(r"[A-Z0-9 ,.-]+", s):
            current = [s, []]
            blocks.append(current)
        elif current:
            current[1].append(s)
        else:
            if blocks:
                raise ValueError("Unexpected text before the first heading")
            blocks.append([None, [s]])
    assert any(name == "START" for name, _ in blocks)
    return title, blocks


def content_markup(blocks):
    parts = []
    for n, (heading, paragraphs) in enumerate(blocks):
        if heading:
            parts.append(f'<section aria-labelledby="section-{n}"><h2 id="section-{n}">{html.escape(heading.title())}</h2>')
        numbered = False
        for paragraph in paragraphs:
            match = re.match(r"^\d+\.\s+(.*)", paragraph)
            if numbered and not match:
                parts.append("</ol>")
                numbered = False
            if match:
                if not numbered:
                    parts.append("<ol>")
                    numbered = True
                parts.append(f"<li>{html.escape(match.group(1))}</li>")
            else:
                parts.append(f"<p>{html.escape(paragraph)}</p>")
        if numbered:
            parts.append("</ol>")
        if heading:
            parts.append("</section>")
    return "\n".join(parts)


def build_html(title, blocks):
    body = content_markup(blocks)
    navigation = "\n".join(
        f'<li><a href="#section-{n}">{html.escape(heading.title())}</a></li>'
        for n, (heading, _) in enumerate(blocks) if heading
    )
    page = f'''<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>{html.escape(title.title())}</title>
<style>body{{font:1.125rem/1.6 system-ui,sans-serif;max-width:48rem;margin:auto;padding:1rem;color:#111;background:white}}a{{color:#0645ad}}h1,h2{{line-height:1.25}}h2{{margin-top:2rem}}li{{margin-block:.4rem}}:focus-visible{{outline:3px solid #0645ad}}</style></head>
<body><header><h1>{html.escape(title.title())}</h1></header>
<nav aria-label="Guide sections"><p>Contents</p><ol>{navigation}</ol></nav>
<main>{body}</main></body></html>'''
    (ROOT / f"{STEM}.html").write_text(page, encoding="utf-8")


def build_word(title, blocks):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(.75)
    section.left_margin = section.right_margin = Inches(.85)
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    for name in ('Title', 'Heading 1', 'Heading 2', 'Normal'):
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    doc.styles["Heading 1"].font.size = Pt(20)
    doc.styles["Heading 2"].font.size = Pt(14)
    doc.add_heading(title.title(), 1)
    for heading, paragraphs in blocks:
        if heading:
            doc.add_heading(heading.title(), 2)
        for p in paragraphs:
            match = re.match(r"^\d+\.\s+(.*)", p)
            doc.add_paragraph(match.group(1) if match else p, style="List Number" if match else "Normal")
    doc.core_properties.title = title.title()
    doc.core_properties.language = "en-US"
    doc.save(ROOT / f"{STEM}.docx")


def build_epub(title, blocks):
    esc = html.escape
    xhtml = f'''<?xml version="1.0" encoding="utf-8"?>
<html xmlns="http://www.w3.org/1999/xhtml" lang="en" xml:lang="en"><head><title>{esc(title.title())}</title></head>
<body><h1>{esc(title.title())}</h1>{content_markup(blocks)}</body></html>'''
    nav_items = "".join(f'<li><a href="content.xhtml#section-{n}">{esc(h.title())}</a></li>' for n, (h, _) in enumerate(blocks) if h)
    nav = f'''<?xml version="1.0" encoding="utf-8"?>
<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops" lang="en" xml:lang="en"><head><title>Contents</title></head>
<body><nav epub:type="toc" id="toc"><h1>Contents</h1><ol>{nav_items}</ol></nav></body></html>'''
    modified = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    opf = f'''<?xml version="1.0" encoding="utf-8"?>
<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="pub-id" xml:lang="en"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/">
<dc:identifier id="pub-id">{UID}</dc:identifier><dc:title>{esc(title.title())}</dc:title><dc:language>en</dc:language><meta property="dcterms:modified">{modified}</meta></metadata>
<manifest><item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/><item id="content" href="content.xhtml" media-type="application/xhtml+xml"/></manifest>
<spine><itemref idref="content"/></spine></package>'''
    container = '''<?xml version="1.0" encoding="utf-8"?><container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles><rootfile full-path="OEBPS/package.opf" media-type="application/oebps-package+xml"/></rootfiles></container>'''
    with ZipFile(ROOT / f"{STEM}.epub", "w") as z:
        z.writestr("mimetype", "application/epub+zip", compress_type=ZIP_STORED)
        for name, data in [("META-INF/container.xml", container), ("OEBPS/package.opf", opf), ("OEBPS/nav.xhtml", nav), ("OEBPS/content.xhtml", xhtml)]:
            z.writestr(name, data.encode("utf-8"), compress_type=ZIP_DEFLATED)


def build_daisy(title, blocks):
    esc = html.escape
    nodes = [("title", title.title(), "doctitle")]
    for i, (heading, paragraphs) in enumerate(blocks):
        if heading:
            nodes.append((f"heading-{i}", heading.title(), "h2"))
        for j, p in enumerate(paragraphs):
            nodes.append((f"text-{i}-{j}", p, "p"))
    front = f'<frontmatter><doctitle>{esc(title.title())}</doctitle></frontmatter>'
    body = [f'<level1 id="guide"><h1 id="title" smilref="guide.smil#text-title">{esc(title.title())}</h1>']
    for i, (heading, paragraphs) in enumerate(blocks):
        if heading:
            body.append(f'<level2 id="section-{i}"><h2 id="heading-{i}" smilref="guide.smil#text-heading-{i}">{esc(heading.title())}</h2>')
        for j, p in enumerate(paragraphs):
            body.append(f'<p id="text-{i}-{j}" smilref="guide.smil#text-text-{i}-{j}">{esc(p)}</p>')
        if heading:
            body.append("</level2>")
    body.append("</level1>")
    dtbook = f'''<?xml version="1.0" encoding="utf-8"?>
<!DOCTYPE dtbook PUBLIC "-//NISO//DTD dtbook 2005-3//EN" "http://www.daisy.org/z3986/2005/dtbook-2005-3.dtd">
<dtbook xmlns="http://www.daisy.org/z3986/2005/dtbook/" version="2005-3" xml:lang="en"><head><meta name="dc:Identifier" content="{UID}"/><meta name="dc:Title" content="{esc(title.title())}"/></head><book>{front}<bodymatter>{''.join(body)}</bodymatter></book></dtbook>'''
    pars = "".join(f'<par id="par-{id_}"><text id="text-{id_}" src="guide.xml#{id_}"/></par>' for id_, _, _ in nodes)
    smil = f'''<?xml version="1.0" encoding="utf-8"?><smil xmlns="http://www.w3.org/2001/SMIL20/"><head><meta name="dtb:uid" content="{UID}"/></head><body><seq id="all">{pars}</seq></body></smil>'''
    points = [f'<navPoint id="nav-title" playOrder="1"><navLabel><text>{esc(title.title())}</text></navLabel><content src="guide.smil#par-title"/></navPoint>']
    for i, (heading, _) in enumerate(blocks):
        if heading:
            points.append(f'<navPoint id="nav-{i}" playOrder="{len(points)+1}"><navLabel><text>{esc(heading.title())}</text></navLabel><content src="guide.smil#par-heading-{i}"/></navPoint>')
    points = [points[0].removesuffix("</navPoint>") + "".join(points[1:]) + "</navPoint>"]
    ncx = f'''<?xml version="1.0" encoding="utf-8"?><ncx xmlns="http://www.daisy.org/z3986/2005/ncx/" version="2005-1" xml:lang="en"><head><meta name="dtb:uid" content="{UID}"/><meta name="dtb:depth" content="2"/><meta name="dtb:totalPageCount" content="0"/><meta name="dtb:maxPageNumber" content="0"/></head><docTitle><text>{esc(title.title())}</text></docTitle><navMap>{''.join(points)}</navMap></ncx>'''
    manifest = '<item id="ncx" href="guide.ncx" media-type="application/x-dtbncx+xml"/><item id="dtbook" href="guide.xml" media-type="application/x-dtbook+xml"/><item id="smil" href="guide.smil" media-type="application/smil"/>'
    opf = f'''<?xml version="1.0" encoding="utf-8"?><package xmlns="http://openebook.org/namespaces/oeb-package/1.0/" unique-identifier="uid"><metadata><dc-metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:Title>{esc(title.title())}</dc:Title><dc:Identifier id="uid">{UID}</dc:Identifier><dc:Language>en</dc:Language></dc-metadata><x-metadata><meta name="dtb:multimediaType" content="textNCX"/><meta name="dtb:multimediaContent" content="text"/></x-metadata></metadata><manifest>{manifest}</manifest><spine toc="ncx"><itemref idref="smil"/></spine></package>'''
    with ZipFile(ROOT / f"{STEM} - DAISY 3.zip", "w", ZIP_DEFLATED) as z:
        for name, data in [("guide.opf", opf), ("guide.xml", dtbook), ("guide.ncx", ncx), ("guide.smil", smil), ("READ ME.txt", "Text-only DAISY 3 guide. Extract all files to one folder and open guide.opf in a DAISY 3 player that supports text-to-speech. No recorded narration is included.\n")]:
            z.writestr(name, data.encode("utf-8"))


def validate(blocks):
    expected = [heading.title() for heading, _ in blocks if heading]
    page = (ROOT / f"{STEM}.html").read_text(encoding="utf-8")
    assert re.findall(r"<h2 id=\"section-\d+\">([^<]+)</h2>", page) == expected
    doc = Document(ROOT / f"{STEM}.docx")
    assert [p.text for p in doc.paragraphs if p.style.name == "Heading 2"] == expected
    with ZipFile(ROOT / f"{STEM}.epub") as z:
        assert z.namelist()[0] == "mimetype" and z.getinfo("mimetype").compress_type == ZIP_STORED
        assert z.read("mimetype") == b"application/epub+zip"
        content = ET.fromstring(z.read("OEBPS/content.xhtml"))
        nav = ET.fromstring(z.read("OEBPS/nav.xhtml"))
        ns = {"x": "http://www.w3.org/1999/xhtml"}
        assert [h.text for h in content.findall(".//x:h2", ns)] == expected
        assert [a.text for a in nav.findall(".//x:a", ns)] == expected
        for name in ("META-INF/container.xml", "OEBPS/package.opf"):
            ET.fromstring(z.read(name))
    with ZipFile(ROOT / f"{STEM} - DAISY 3.zip") as z:
        dt = ET.fromstring(z.read("guide.xml"))
        ncx = ET.fromstring(z.read("guide.ncx"))
        smil = ET.fromstring(z.read("guide.smil"))
        ET.fromstring(z.read("guide.opf"))
        dn = {"d": "http://www.daisy.org/z3986/2005/dtbook/"}
        nn = {"n": "http://www.daisy.org/z3986/2005/ncx/"}
        sn = {"s": "http://www.w3.org/2001/SMIL20/"}
        assert [h.text for h in dt.findall(".//d:h2", dn)] == expected
        assert [h.text for h in ncx.findall(".//n:navPoint/n:navLabel/n:text", nn)][1:] == expected
        ids = {e.get("id") for e in dt.iter() if e.get("id")}
        assert all(e.get("src").split("#")[1] in ids for e in smil.findall(".//s:text", sn))
        smil_ids = {e.get("id") for e in smil.iter() if e.get("id")}
        assert all(e.get("smilref").split("#")[1] in smil_ids for e in dt.iter() if e.get("smilref"))


def main():
    from release_documents import load_notes
    version=load_notes(ROOT)['version']
    source=SOURCE.read_text(encoding='utf-8-sig')
    source=re.sub(r'Guide build: [^ ]+\.', 'Guide build: '+version+'.', source, count=1)
    SOURCE.write_text(source,encoding='utf-8')
    title, blocks = sections()
    build_html(title, blocks)
    build_word(title, blocks)
    build_epub(title, blocks)
    build_daisy(title, blocks)
    validate(blocks)
    from braille_documents import build_brf
    build_brf(SOURCE, ROOT / f"{STEM}.brf")
    from release_documents import build
    build(ROOT)
    from guide_files import write_manifest
    from release_documents import load_notes
    write_manifest(ROOT, load_notes(ROOT)['version'])
    print("Validated six manual formats and six release-note formats, including UEB BRF.")


if __name__ == "__main__":
    main()

