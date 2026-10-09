"""Verify spoken formatting survives a DOCX save, reopen, and undo."""
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET

from document_editor import VoiceDocument
from docx import Document


class FormattingTests(unittest.TestCase):
    def test_short_spoken_spacing_and_prompt_do_not_dictate_answers(self):
        with tempfile.TemporaryDirectory() as folder:
            doc = VoiceDocument(Path(folder))
            self.assertIn('Write a paragraph first', doc.process('custom spacing'))
            doc.append_text('Keep this sentence.')
            self.assertIn('1.00 line spacing', doc.process('single space'))
            self.assertIn('2.00 line spacing', doc.process('double space'))
            self.assertIn('What spacing', doc.process('custom spacing'))
            self.assertTrue(doc.pending_spacing)
            before = doc.paragraphs[0].line_spacing_twips
            self.assertIn('value first', doc.process('okay'))
            self.assertIn('Say another value', doc.process('30'))
            self.assertEqual(doc.paragraphs[0].text, 'Keep this sentence.')
            self.assertIn('I heard 1.75 line spacing', doc.process('one point seven five'))
            self.assertEqual(doc.paragraphs[0].line_spacing_twips, before)
            self.assertIn('Set 1.75 line spacing', doc.process('confirm that'))
            self.assertFalse(doc.pending_spacing)
            self.assertIn('What spacing', doc.process('custom spacing'))
            self.assertIn('I heard 1.50', doc.process('one point five'))
            self.assertIn('Canceled', doc.process('cancel'))
            self.assertIn('1.75', doc.process('read line spacing'))
            self.assertIn('What spacing', doc.process('custom spacing'))
            self.assertIn('I heard exactly 18 points', doc.process('exactly eighteen points'))
            self.assertIn('Set exactly 18 points', doc.process('okay'))
            self.assertEqual(doc.paragraphs[0].text, 'Keep this sentence.')
            reopened = VoiceDocument.open_existing(Path(folder), doc.title)
            self.assertIn('exactly 18 points', reopened.process('read line spacing'))
            self.assertIn('Undid', doc.process('undo'))
            self.assertIn('1.75 line spacing', doc.process('read line spacing'))

    def test_standard_custom_and_exact_line_spacing_round_trip(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            doc = VoiceDocument(root)
            doc.append_text('First paragraph')
            self.assertIn('1.75 line spacing', doc.process('line spacing to one point seven five'))
            self.assertIn('1.75', doc.process('read line spacing'))
            doc.process('new paragraph')
            doc.append_text('Second paragraph')
            self.assertIn('1.75', doc.process('read line spacing'))
            self.assertIn('exactly 18 points', doc.process('line spacing exactly eighteen points'))
            doc.process('new paragraph')
            doc.append_text('Third paragraph')
            self.assertIn('at least 21.5 points', doc.process('line spacing at least 21.5 points'))
            self.assertIn('Choose line spacing', doc.process('line spacing 0.2'))
            self.assertIn('4 to 144', doc.process('line spacing exactly 200 points'))
            self.assertIn('at least 21.5 points', doc.process('read line spacing'))
            with ZipFile(doc.path) as archive:
                xml = ET.fromstring(archive.read('word/document.xml'))
            ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
            attrs = [p.attrib for p in xml.findall('.//w:spacing', ns)]
            key = '{' + ns['w'] + '}'
            self.assertEqual([(a[key+'line'], a[key+'lineRule']) for a in attrs],
                             [('420','auto'), ('360','exact'), ('430','atLeast')])
            word = Document(doc.path)
            self.assertEqual(word.paragraphs[0].paragraph_format.line_spacing, 1.75)
            self.assertEqual(word.paragraphs[1].paragraph_format.line_spacing.pt, 18)
            reopened = VoiceDocument.open_existing(root, doc.title)
            reopened.cursor = 2
            self.assertIn('at least 21.5', reopened.process('read line spacing'))
            self.assertIn('Restored default', reopened.process('reset line spacing'))
            self.assertIn('default line spacing', reopened.process('read line spacing'))
            reopened.process('undo')
            self.assertIn('at least 21.5', reopened.process('read line spacing'))
            self.assertIn('2.50 line spacing', reopened.process('line spacing two and a half'))
            self.assertIn('exactly 120 points', reopened.process('line spacing exactly one hundred and twenty points'))
            self.assertIn('144 points', reopened.process('line spacing exactly 144 points'))
            self.assertIn('10.00 line spacing', reopened.process('line spacing 10'))
            self.assertIn('throughout the document', reopened.process('set document line spacing to 1.25'))
            self.assertIn('1.25 line spacing throughout', reopened.process('read document line spacing'))
            again = VoiceDocument.open_existing(root, doc.title)
            self.assertTrue(all(p.line_spacing_twips == 300 for p in again.paragraphs))

    def test_word_formatting_survives_later_dictation(self):
        with tempfile.TemporaryDirectory() as folder:
            doc = VoiceDocument(Path(folder))
            doc.append_text('Alpha beta gamma.')
            self.assertIn('beta', doc.process('next word'))
            doc.process('bold word')
            doc.process('underline word')
            doc.process('set word font to Arial')
            doc.append_text('Delta.')
            restored = VoiceDocument.open_existing(Path(folder), doc.title)
            runs = restored.paragraphs[0].runs
            self.assertEqual(''.join(r.text for r in runs), 'Alpha beta gamma. Delta.')
            marked = next(r for r in runs if r.text == 'beta')
            self.assertTrue(marked.bold and marked.underline)
            self.assertEqual(marked.font, 'Arial')
            self.assertFalse(runs[-1].bold)

    def test_formatting_and_table_round_trip(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            doc = VoiceDocument(root)
            doc.append_text('Introduction')
            for command in ('heading level two', 'make this bold', 'make this italic',
                            'underline this', 'set font to Arial', 'set size to 14 points',
                            'align center'):
                doc.process(command)
            self.assertIn('Heading 2', doc.process('read formatting'))
            doc.process('new paragraph')
            doc.append_text('An item')
            doc.process('bulleted list')
            doc.process('new paragraph')
            doc.append_text('A second item')
            doc.process('numbered list')
            doc.process('create table with 2 rows and 2 columns')
            doc.append_text('Top left')
            doc.process('next cell')
            doc.append_text('Top right')
            doc.process('next row')
            doc.append_text('Bottom right')
            doc.process('leave table')
            doc.append_text('After the table')
            with ZipFile(doc.path) as archive:
                xml = archive.read('word/document.xml')
                numbering = archive.read('word/numbering.xml')
            ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
            tree = ET.fromstring(xml)
            self.assertEqual(len(tree.findall('.//w:tbl/w:tr', ns)), 2)
            self.assertEqual(len(tree.findall('.//w:tbl/w:tr/w:tc', ns)), 4)
            self.assertEqual(len(tree.findall('.//w:numPr', ns)), 2)
            self.assertIn(b'Heading2', xml)
            self.assertIn(b'w:u', xml)
            self.assertIn(b'Arial', xml)
            self.assertIn(b'bullet', numbering)
            restored = VoiceDocument.open_existing(root, doc.title)
            self.assertEqual(restored.paragraphs[0].kind, 'Heading 2')
            self.assertTrue(restored.paragraphs[0].underline)
            self.assertEqual(restored.paragraphs[0].font, 'Arial')
            self.assertEqual(restored.paragraphs[0].size, 14)
            self.assertEqual(restored.paragraphs[0].alignment, 'center')
            self.assertEqual([p.kind for p in restored.paragraphs[1:3]], ['Bulleted list', 'Numbered list'])
            restored.cursor = 3
            self.assertIn('Row 1, column 1', restored.process('read cell'))
            self.assertIn('Row 1, column 2', restored.process('next cell'))
            self.assertIn('Row 2, column 2', restored.process('next row'))
            self.assertIn('After the table', restored.paragraphs[-1].text)

    def test_undo_restores_table_and_rejects_unsafe_dimensions(self):
        with tempfile.TemporaryDirectory() as folder:
            doc = VoiceDocument(Path(folder))
            self.assertIn('choose up to', doc.process('create table with 300 rows and 2 columns'))
            self.assertEqual(doc.paragraphs, [])
            doc.process('create table with 2 rows and 2 columns')
            doc.append_text('Value')
            doc.process('delete table')
            self.assertEqual(doc.paragraphs, [])
            doc.process('undo')
            self.assertEqual(len(doc.paragraphs), 4)
            self.assertEqual(doc.paragraphs[0].text, 'Value')

    def test_add_row_and_column_preserve_rectangular_table(self):
        with tempfile.TemporaryDirectory() as folder:
            doc = VoiceDocument(Path(folder))
            doc.process('create table with two rows and two columns')
            doc.append_text('First')
            self.assertIn('Row 2, column 1', doc.process('add row'))
            self.assertIn('Row 2, column 2', doc.process('add column'))
            doc.append_text('New cell')
            reopened = VoiceDocument.open_existing(Path(folder), doc.title)
            cells = [p for p in reopened.paragraphs if p.table_id]
            self.assertEqual(len(cells), 9)
            self.assertEqual({(p.row,p.column) for p in cells},
                             {(r,c) for r in range(1,4) for c in range(1,4)})
            self.assertEqual(next(p.text for p in cells if (p.row,p.column)==(2,2)), 'New cell')
            self.assertIn('Deleted the column', doc.process('delete column'))
            self.assertIn('Deleted the row', doc.process('delete row'))
            reread = VoiceDocument.open_existing(Path(folder), doc.title)
            self.assertEqual(len([p for p in reread.paragraphs if p.table_id]), 4)


if __name__ == '__main__':
    unittest.main()
