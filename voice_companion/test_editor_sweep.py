"""Word processor, Notepad, reader and File manager: bugs a new user met (sweep, October 29, 2026)."""
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import companion as app
from document_editor import Paragraph, VoiceDocument
from test_document_save_as import DocumentFlowCase


class FailedSaveLeavesNoStrayFile(DocumentFlowCase):
    def test_a_locked_file_leaves_no_extra_document_in_the_list(self):
        made = VoiceDocument(app.documents_folder(), title='Letter')
        made.paragraphs = [Paragraph('Dear Bob.')]; made.save()
        with patch.object(Path, 'replace', side_effect=PermissionError('open in Word')):
            with self.assertRaises(OSError):
                made.save()
        names = sorted(p.name for p in app.documents_folder().glob('*.docx'))
        self.assertEqual(names, ['Letter.docx'])
        app.handle('list documents', 'word_processor')
        self.assertIn('1 of 1', self.last())


class SpokenRenameInFileManager(DocumentFlowCase):
    """Rename by voice: recognizers add a period, and people say "dot" for the dot."""
    def setUp(self):
        super().setUp()
        from file_manager import FileManager
        from test_file_manager import FakeFiles, settle
        self.settle = settle
        self.disk = self.root / 'Disk'; self.disk.mkdir()
        (self.disk / 'notes.txt').write_text('Milk.')
        self.fm = FileManager(FakeFiles(self.disk))
        self.stack.enter_context(patch.object(app, 'FILE_MANAGER', self.fm))
        self.fm.refresh(str(self.disk)); settle(self.fm)
        self.fm.choose('notes.txt')
        self.assertIn('Rename notes.txt', self.fm.command('rename'))

    def renamed_to(self, heard, typed=False):
        app.handle(heard, 'files', typed=typed); self.settle(self.fm)
        actions = self.fm.backend.actions
        return actions[-1][3] if actions and actions[-1][0] == 'rename' else None

    def test_a_period_the_recognizer_added_is_left_off(self):
        self.assertEqual(self.renamed_to('Shopping.txt.'), 'Shopping.txt')

    def test_dot_and_the_file_type_said_as_words(self):
        for heard, name in (('shopping dot txt', 'shopping.txt'), ('shopping dot t x t', 'shopping.txt'),
                            ('Shopping list dot docx.', 'Shopping list.docx')):
            with self.subTest(heard=heard):
                self.fm.backend.actions.clear()
                self.fm.pending = ('rename', [str(self.disk / 'notes.txt')])
                self.assertEqual(self.renamed_to(heard), name)

    def test_a_typed_name_is_used_as_typed(self):
        self.assertEqual(self.renamed_to('shopping dot txt', typed=True), 'shopping dot txt')

    def test_a_typed_name_ending_in_a_period_says_why_it_cannot_be_used(self):
        self.assertIsNone(self.renamed_to('Shopping.', typed=True))
        self.assertIn('cannot end with a period or a space', self.last())


class ComplexWordFileOpensInTheReader(DocumentFlowCase):
    def test_the_reason_is_said_in_plain_words(self):
        from docx import Document
        made = Document(); made.add_paragraph('Body text.'); made.sections[0].header.paragraphs[0].text = 'Top'
        path = self.elsewhere / 'Complex.docx'; made.save(path)
        self.assertEqual(app.task_navigation().open_document(path), 'file_reader')
        self.assertIn('Opened Complex.docx. This Word file has features the editor cannot keep, '
                      'so it opened in the Document reader. You can read it, but not change it.', self.last())


class OpenDocumentByName(DocumentFlowCase):
    """"Open document" followed by a name does what choosing the file in Open a document does."""
    def setUp(self):
        super().setUp(); self.default.mkdir(parents=True, exist_ok=True)

    def test_a_complex_word_file_opens_in_the_reader(self):
        from docx import Document
        made = Document(); made.add_paragraph('Body text.'); made.sections[0].header.paragraphs[0].text = 'Top'
        made.save(self.default / 'Complex.docx')
        self.assertEqual(app.handle('Open document Complex.', 'word_processor'), 'file_reader')
        self.assertIn('features the editor cannot keep', self.last())
        self.assertNotIn('test editor', ' '.join(self.said))
        self.assertEqual(self.press('Escape', 'file_reader'), 'word_processor')
        self.assertTrue(self.last().startswith('Word processor menu.'))

    def test_a_damaged_file_is_named_in_plain_words(self):
        (self.default / 'Broken.docx').write_bytes(b'PK not really')
        self.assertEqual(app.handle('open document Broken', 'word_processor'), 'word_processor')
        self.assertEqual(self.last(), 'This document is empty or damaged, so it could not be opened.')


class ListsOverAMenu(DocumentFlowCase):
    """A list opened over the Word processor or Notepad menu owns the keys, as its prompt says."""
    def setUp(self):
        super().setUp(); self.default.mkdir(parents=True, exist_ok=True)
        for name in ('Alpha', 'Bravo'):
            made = VoiceDocument(self.default, title=name); made.paragraphs = [Paragraph(name + ' text.')]; made.save()

    def test_enter_on_list_documents_opens_the_document(self):
        mode = app.handle('list documents', 'word_processor')
        mode = self.press('Down', mode)
        self.assertIn('Bravo', self.last())
        self.assertEqual(self.press('Enter', mode), 'document')
        self.assertEqual(app.document.title, 'Bravo')

    def test_home_and_end_move_in_the_list(self):
        mode = app.handle('list documents', 'word_processor')
        mode = self.press('End', mode)
        self.assertIn('Bravo', self.last())
        mode = self.press('Home', mode)
        self.assertIn('Alpha', self.last())
        self.assertEqual(self.press('Enter', mode), 'document')
        self.assertEqual(app.document.title, 'Alpha')

    def test_escape_closes_the_list_and_stays_on_the_menu(self):
        mode = app.handle('list documents', 'word_processor')
        self.assertEqual(self.press('Escape', mode), 'word_processor')
        self.assertIsNone(app.NAMED_LIST)

    def test_enter_on_the_notes_list_opens_the_note(self):
        import notepad
        notepad.add(app.APP, 'Buy milk.')
        mode = self.press('Ctrl+O', app.handle('notepad', 'awake'))
        self.assertIn('Buy milk', self.last())
        self.assertEqual(self.press('Enter', mode), 'note')
        self.assertIn('Buy milk.', ' '.join(self.said[-2:]))


class SpokenDocumentNames(DocumentFlowCase):
    """People say the file type: "My letter dot docx" is My letter.docx, not My letter dot docx.docx."""
    def test_new_document(self):
        app.handle('create a document', 'awake'); app.handle('My letter dot docx.', 'document_name')
        self.assertEqual(app.handle('confirm', 'document_name'), 'document')
        self.assertEqual(app.document.path.name, 'My letter.docx')

    def test_rename_and_save_as(self):
        app.document = VoiceDocument(self.default, title='First'); app.document.append_text('Words.')
        app.handle('name document Second dot d o c x', 'document')
        self.assertEqual(app.document.path.name, 'Second.docx')
        app.handle('save as Third dot docx', 'document'); app.handle('confirm', 'document_name')
        self.assertEqual(app.document.path.name, 'Third.docx')


class OpenFromADocumentThenEscape(DocumentFlowCase):
    """Control+O in a document, then Escape: back to the document, as when a new document is canceled."""
    def setUp(self):
        super().setUp()
        import file_manager, regular_browser
        from test_file_manager import FakeFiles, settle
        self.settle = settle
        disk = self.root / 'Disk'; disk.mkdir()
        self.stack.enter_context(patch.object(file_manager, 'ShellFiles', lambda: FakeFiles(disk)))
        self.stack.enter_context(patch.object(regular_browser, 'data_folder', lambda: self.root))

    def escape_out(self, mode):
        for _ in range(5):
            while app.FILE_MANAGER.busy:
                import time; time.sleep(.01); mode = app.poll_file_manager(mode)
            if mode != 'files': return mode
            mode = self.press('Escape', mode)
        return mode

    def test_escape_from_open_returns_to_the_document(self):
        app.document = VoiceDocument(self.default, title='Letter'); app.document.append_text('Dear Bob.')
        mode = self.press('Ctrl+O', 'document')
        self.assertEqual(mode, 'files')
        self.assertEqual(self.escape_out(mode), 'document')
        self.assertEqual(self.last(), 'Canceled. Back to document.')
        self.assertEqual(app.document.title, 'Letter')

    def test_escape_from_open_on_the_menu_still_returns_to_the_menu(self):
        mode = self.press('Ctrl+O', 'word_processor')
        self.assertEqual(self.escape_out(mode), 'word_processor')


class EmptyFolderKeys(DocumentFlowCase):
    """In an empty folder, every key that needs an item says the folder is empty."""
    def setUp(self):
        super().setUp()
        from file_manager import FileManager
        from test_file_manager import FakeFiles, settle
        empty = self.root / 'Empty'; empty.mkdir()
        self.fm = FileManager(FakeFiles(empty))
        self.stack.enter_context(patch.object(app, 'FILE_MANAGER', self.fm))
        self.fm.refresh(str(empty)); settle(self.fm)

    def test_keys_that_need_an_item(self):
        for key in ('Space', 'Ctrl+A', 'Delete', 'F2', 'Ctrl+C', 'Ctrl+X', 'Letter:a'):
            with self.subTest(key=key):
                self.said.clear()
                self.assertEqual(self.press(key, 'files'), 'files')
                self.assertEqual(self.said, ['This folder is empty.'])


class CopyOrCutWhatIsSelected(DocumentFlowCase):
    """With text selected, "Copy" or "Cut" alone acts on it, as Control+C and Control+X do.
    With nothing selected, the word is still written."""
    def setUp(self):
        super().setUp()
        import document_editor
        self.clip = []
        self.stack.enter_context(patch.object(document_editor, 'copy_to_clipboard', self.clip.append))
        app.document = VoiceDocument(self.default, title='Sel'); app.document.append_text('The quick brown fox.')

    def text(self):
        return ' '.join(p.text for p in app.document.paragraphs)

    def test_copy(self):
        app.handle('select all', 'document'); app.handle('Copy.', 'document')
        self.assertEqual(self.text(), 'The quick brown fox.')
        self.assertEqual(self.clip, ['The quick brown fox.'])

    def test_cut(self):
        app.handle('select all', 'document'); app.handle('cut', 'document')
        self.assertEqual(self.text(), '')
        self.assertEqual(self.clip, ['The quick brown fox.'])
        app.handle('undo', 'document')
        self.assertEqual(self.text(), 'The quick brown fox.')

    def test_without_a_selection_the_word_is_text(self):
        app.handle('go to end of document', 'document'); app.handle('Copy', 'document')
        self.assertIn('Copy', self.text())
        self.assertEqual(self.clip, [])


class MovingPastTheEnds(DocumentFlowCase):
    """Next word, sentence or line at the end says so, as Next paragraph does; a repeated word sounded like a move."""
    def test_document(self):
        app.document = VoiceDocument(self.default, title='Ends')
        app.document.paragraphs = [Paragraph('First one.'), Paragraph('Second one.')]
        app.handle('go to end of document', 'document')
        for unit in ('word', 'sentence', 'line'):
            with self.subTest(unit=unit):
                app.handle('go to end of document', 'document')
                app.handle('next ' + unit, 'document')
                self.assertTrue(self.last().startswith('End of document. '), self.last())
                app.handle('go to beginning of document', 'document')
                app.handle('previous ' + unit, 'document')
                self.assertTrue(self.last().startswith('Beginning of document. '), self.last())
                app.handle('next ' + unit, 'document')
                self.assertFalse(self.last().startswith(('End of', 'Beginning of')), self.last())

    def test_document_reader(self):
        from document_reader import DocumentReader
        reader = DocumentReader('notes.txt', 'Dear Bob. See you soon.')
        self.stack.enter_context(patch.object(app, 'FILE_READER', reader))
        app.handle('go to end of document', 'file_reader')
        app.handle('next word', 'file_reader')
        self.assertEqual(self.last(), 'End of document. soon.')
        app.handle('go to beginning of document', 'file_reader')
        app.handle('previous sentence', 'file_reader')
        self.assertEqual(self.last(), 'Beginning of document. Dear Bob.')

    def test_note(self):
        app.note_buffer[:] = ['Buy milk. Call Bob.']
        app.handle('next sentence', 'note')
        app.handle('next sentence', 'note')
        self.assertEqual(self.last(), 'End of note. Call Bob.')


if __name__ == '__main__':
    unittest.main()
